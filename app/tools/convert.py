"""Conversions: PDF<->images, Word/PPT/Excel, Markdown, HTML."""
from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import time
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageFilter, ImageOps

from .. import config
from . import netguard
from .base import (Ctx, ToolError, guess_mime, open_pdf, parse_pages, run_cmd, safe_stem, save_pdf, tool,
                   zip_outputs)
from .pages import PAGE_SIZES, pdf_of

try:  # HEIC/HEIF support if available
    import pi_heif
    pi_heif.register_heif_opener()
except Exception:  # pragma: no cover
    pass

IMG_EXT = ("jpg", "jpeg", "png", "webp", "gif", "bmp", "tif", "tiff", "heic", "heif")
Image.MAX_IMAGE_PIXELS = 120_000_000


# ---------------------------------------------------------------------------
@tool("pdf2jpg", timeout=600)
def pdf2jpg(ctx: Ctx):
    doc = pdf_of(ctx, 0)
    mode = ctx.opt("mode", "pages")
    fmt = ctx.opt("format", "jpg")
    files: list[tuple[Path, str]] = []
    stem = ctx.stem()
    if mode == "images":
        seen = set()
        for pno in range(len(doc)):
            for img in doc[pno].get_images(full=True):
                xref = img[0]
                if xref in seen:
                    continue
                seen.add(xref)
                try:
                    info = doc.extract_image(xref)
                except Exception:
                    continue
                if min(info.get("width", 0), info.get("height", 0)) < 32:
                    continue
                ext = info["ext"] if info["ext"] in ("jpeg", "jpg", "png") else "png"
                data = info["image"]
                if info["ext"] not in ("jpeg", "jpg", "png"):
                    try:
                        pil = Image.open(io.BytesIO(data)).convert("RGBA" if "A" in Image.open(io.BytesIO(data)).mode else "RGB")
                        buf = io.BytesIO()
                        pil.save(buf, "PNG")
                        data = buf.getvalue()
                    except Exception:
                        continue
                p = ctx.tmp / f"img_{len(files):04d}.{ext}"
                p.write_bytes(data)
                files.append((p, f"{stem}_p{pno + 1}_img{len(files) + 1}.{'jpg' if ext == 'jpeg' else ext}"))
        if not files:
            raise ToolError("no_images", "No embedded images were found in this PDF.")
    else:
        dpi = max(36, min(300, int(ctx.opt("dpi", 150))))
        quality = max(30, min(100, int(ctx.opt("quality", 88))))
        pages = parse_pages(ctx.opt("pages", "all"), len(doc))
        if len(pages) > 300:
            raise ToolError("too_many", "Please select at most 300 pages at a time.")
        for n, p in enumerate(pages):
            pg = doc[p]
            d = dpi
            px = pg.rect.width * pg.rect.height * (d / 72) ** 2
            if px > 60e6:
                d = int(d * (60e6 / px) ** 0.5)
            pix = pg.get_pixmap(dpi=d, alpha=False)
            f = ctx.tmp / f"page_{p + 1:04d}.{fmt}"
            if fmt == "png":
                pix.save(str(f))
            else:
                f.write_bytes(pix.tobytes("jpeg", jpg_quality=quality))
            files.append((f, f"{stem}_page_{p + 1}.{fmt}"))
            ctx.progress(n / len(pages))
    ctx.data["count"] = len(files)
    if len(files) == 1:
        ctx.add_output(files[0][0], files[0][1])
    else:
        zip_outputs(ctx, files, f"{stem}_images.zip")


# ---------------------------------------------------------------------------
def _prepare_image(path: Path, flt: str, tmp: Path, idx: int):
    """Return (bytes, (w,h)px, dpi). Applies EXIF rotation and scan filters."""
    img = Image.open(path)
    img.load()
    dpi = img.info.get("dpi", (0, 0))
    dpi = float(dpi[0]) if dpi and dpi[0] and dpi[0] > 20 else 0.0
    fmt = (img.format or "").upper()
    exif_rot = False
    try:
        exif_rot = (img.getexif().get(274, 1) or 1) != 1
    except Exception:
        pass
    img = ImageOps.exif_transpose(img)
    if max(img.size) > 5000:
        img.thumbnail((5000, 5000), Image.LANCZOS)
        fmt = ""
    if flt in ("gray", "bw", "enhance"):
        g = ImageOps.autocontrast(ImageOps.grayscale(img), cutoff=1)
        if flt == "enhance":
            g = ImageOps.autocontrast(g.filter(ImageFilter.UnsharpMask(radius=2, percent=120)), cutoff=2)
        if flt == "bw":
            # local-mean threshold keeps text readable on uneven lighting
            blur = g.filter(ImageFilter.GaussianBlur(radius=max(8, min(g.size) // 40)))
            g = _bw_mask(g, blur)
        img, fmt = g, ""
    buf = io.BytesIO()
    if fmt == "JPEG" and not exif_rot and flt == "none":
        return path.read_bytes(), img.size, dpi
    if img.mode in ("RGBA", "LA", "P") and (img.mode != "P" or "transparency" in img.info):
        img = img.convert("RGBA")
        bg = Image.new("RGB", img.size, (255, 255, 255))
        bg.paste(img, mask=img.getchannel("A"))
        img = bg
    img = img.convert("L" if img.mode == "L" else "RGB")
    img.save(buf, "JPEG", quality=90)
    return buf.getvalue(), img.size, dpi


def _bw_mask(g: Image.Image, blur: Image.Image) -> Image.Image:
    import numpy as np

    a = np.asarray(g, dtype=np.int16)
    b = np.asarray(blur, dtype=np.int16)
    return Image.fromarray(np.where(a > b - 12, 255, 0).astype("uint8"), "L")


def _images_to_pdf(ctx: Ctx, flt="none", default_filter="none"):
    flt = ctx.opt("filter", default_filter if flt == "none" else flt)
    size_opt = ctx.opt("page_size", "fit")
    orient = ctx.opt("orientation", "auto")
    margin = {"none": 0.0, "small": 10.0, "big": 20.0}.get(str(ctx.opt("margin", "none")), 0.0) * 72 / 25.4
    out = fitz.open()
    for i, p in enumerate(ctx.paths):
        try:
            data, (w, h), dpi = _prepare_image(p, flt, ctx.tmp, i)
        except Exception:
            raise ToolError("bad_image", f"{ctx.names[i]} could not be read as an image.")
        if size_opt == "fit":
            k = 72 / (dpi or 96)
            pw, ph = w * k + 2 * margin, h * k + 2 * margin
            scale = min(1.0, 1500 / max(pw, ph))
            pw, ph = pw * scale, ph * scale
            margin_eff = margin * scale
        else:
            pw, ph = PAGE_SIZES.get(size_opt, PAGE_SIZES["a4"])
            land = (orient == "landscape") or (orient == "auto" and w > h)
            if land:
                pw, ph = max(pw, ph), min(pw, ph)
            else:
                pw, ph = min(pw, ph), max(pw, ph)
            margin_eff = margin
        pg = out.new_page(width=pw, height=ph)
        pg.insert_image(fitz.Rect(margin_eff, margin_eff, pw - margin_eff, ph - margin_eff), stream=data,
                        keep_proportion=True)
        ctx.progress(i / len(ctx.paths))
    return out


@tool("jpg2pdf", accept=IMG_EXT, multi=True, min_files=1, max_files=100, timeout=600)
def jpg2pdf(ctx: Ctx):
    out = _images_to_pdf(ctx)
    if ctx.opt("merge", True) or len(out) == 1:
        dst = ctx.new_path("images.pdf")
        save_pdf(out, dst)
        ctx.add_output(dst, f"{Path(ctx.names[0]).stem if len(ctx.names) == 1 else 'images'}.pdf")
    else:
        files = []
        for i in range(len(out)):
            one = fitz.open()
            one.insert_pdf(out, from_page=i, to_page=i)
            f = ctx.tmp / f"i{i}.pdf"
            save_pdf(one, f)
            files.append((f, f"{Path(ctx.names[i]).stem}.pdf"))
        zip_outputs(ctx, files, "images_pdf.zip")


@tool("scan", accept=IMG_EXT, multi=True, min_files=1, max_files=100, timeout=600)
def scan(ctx: Ctx):
    out = _images_to_pdf(ctx, default_filter="enhance")
    dst = ctx.new_path("scan.pdf")
    save_pdf(out, dst)
    ctx.add_output(dst, f"scan_{time.strftime('%Y%m%d_%H%M')}.pdf")


# ---------------------------------------------------------------------------
def _has_text(doc) -> bool:
    for i in range(min(len(doc), 10)):
        if len(doc[i].get_text("text").strip()) > 20:
            return True
    return False


def _maybe_ocr(ctx: Ctx, src: Path) -> Path:
    """Scanned input: optionally OCR first so conversions have real text."""
    doc = open_pdf(src, ctx.opts.get("password"))
    scanned = not _has_text(doc)
    doc.close()
    if not scanned:
        return src
    if not ctx.opts.get("ocr"):
        raise ToolError("scanned", "This PDF looks scanned (no text layer). Enable OCR and try again.")
    from .optimize import ocr_file
    ctx.progress(0.05, "OCR")
    dst = ctx.tmp / "ocr_in.pdf"
    ocr_file(ctx, src, dst, langs=ctx.opt("ocr_langs", "eng+tha"), mode="skip")
    return dst


@tool("pdf2word", timeout=900, mem_gb=6)
def pdf2word(ctx: Ctx):
    from pdf2docx import Converter

    src = ctx.paths[0]
    pw = ctx.opts.get("password")
    if pw:
        d = open_pdf(src, pw)
        src = ctx.tmp / "dec.pdf"
        d.save(str(src))
    src = _maybe_ocr(ctx, src)
    dst = ctx.new_path("converted.docx")
    cv = Converter(str(src))
    try:
        total = len(cv.fitz_doc)
        pages = parse_pages(ctx.opt("pages", "all"), total)
        cv.convert(str(dst), start=pages[0], end=pages[-1] + 1, multi_processing=False)
    finally:
        cv.close()
    ctx.add_output(dst, f"{ctx.stem()}.docx")


@tool("pdf2pptx", timeout=900, mem_gb=6)
def pdf2pptx(ctx: Ctx):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    doc = pdf_of(ctx, 0)
    mode = ctx.opt("mode", "editable")
    pages = parse_pages(ctx.opt("pages", "all"), len(doc))
    prs = Presentation()
    first = doc[pages[0]].rect
    k = min(1.0, 3600 / max(first.width, first.height))  # keep within PowerPoint limits
    prs.slide_width, prs.slide_height = Emu(int(first.width * 12700 * k)), Emu(int(first.height * 12700 * k))
    blank = prs.slide_layouts[6]
    for n, p in enumerate(pages):
        ctx.progress(n / len(pages), f"{n + 1}/{len(pages)}")
        slide = prs.slides.add_slide(blank)
        pg = doc[p]
        sx = prs.slide_width / pg.rect.width
        sy = prs.slide_height / pg.rect.height
        scanned = any(
            (fitz.Rect(b["bbox"]) & pg.rect).get_area() > 0.85 * pg.rect.get_area()
            for b in pg.get_text("dict")["blocks"] if b["type"] == 1)
        if mode == "image" or scanned:
            pix = pg.get_pixmap(dpi=150, alpha=False)
            slide.shapes.add_picture(io.BytesIO(pix.tobytes("jpeg", jpg_quality=88)), 0, 0,
                                     prs.slide_width, prs.slide_height)
            continue
        tmp = fitz.open()
        tmp.insert_pdf(doc, from_page=p, to_page=p)
        tp = tmp[0]
        blocks = [b for b in tp.get_text("dict")["blocks"] if b["type"] == 0]
        for b in blocks:
            tp.add_redact_annot(fitz.Rect(b["bbox"]), fill=False)
        if blocks:
            tp.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
        pix = tp.get_pixmap(dpi=150, alpha=False)
        slide.shapes.add_picture(io.BytesIO(pix.tobytes("jpeg", jpg_quality=90)), 0, 0,
                                 prs.slide_width, prs.slide_height)
        for b in blocks:
            for line in b["lines"]:
                spans = [s for s in line["spans"] if s["text"].strip() or len(line["spans"]) == 1]
                if not spans:
                    continue
                x0, y0, x1, y1 = line["bbox"]
                tb = slide.shapes.add_textbox(Emu(int(x0 * sx)), Emu(int(y0 * sy)),
                                              Emu(int((x1 - x0) * sx * 1.08) + 20000), Emu(int((y1 - y0) * sy)))
                tf = tb.text_frame
                tf.word_wrap = False
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                para = tf.paragraphs[0]
                for s in spans:
                    run = para.add_run()
                    run.text = s["text"]
                    f = run.font
                    f.size = Pt(max(4, s["size"] * k * 0.98))
                    f.bold = bool(s["flags"] & 16)
                    f.italic = bool(s["flags"] & 2)
                    name = re.sub(r"^[A-Z]{6}\+", "", s["font"]).split("-")[0].split(",")[0]
                    f.name = name or "Calibri"
                    c = s["color"]
                    f.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
    dst = ctx.new_path("converted.pptx")
    prs.save(str(dst))
    ctx.add_output(dst, f"{ctx.stem()}.pptx")


# ---------------------------------------------------------------------------
_NUM = re.compile(r"^-?\(?\d{1,3}(,\d{3})*(\.\d+)?\)?%?$|^-?\d+(\.\d+)?$")


def _cell(v):
    if v is None:
        return ""
    s = str(v).strip().replace("\n", " ")
    if _NUM.match(s) and not (len(s) > 1 and s[0] == "0" and s[1].isdigit()) and not s.endswith("%") \
            and "(" not in s:
        try:
            n = float(s.replace(",", ""))
            return int(n) if n.is_integer() and "." not in s else n
        except ValueError:
            return s
    return s


@tool("pdf2excel", timeout=900, mem_gb=6)
def pdf2excel(ctx: Ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    src = ctx.paths[0]
    pw = ctx.opts.get("password")
    if pw:
        d = open_pdf(src, pw)
        src = ctx.tmp / "dec.pdf"
        d.save(str(src))
    src = _maybe_ocr(ctx, src)
    doc = open_pdf(src)
    pages = parse_pages(ctx.opt("pages", "all"), len(doc))
    layout = ctx.opt("layout", "table")
    strategy = ctx.opt("strategy", "lines")
    wb = Workbook()
    wb.remove(wb.active)
    n_tables = 0
    single = wb.create_sheet("Data") if layout == "single" else None
    for k, p in enumerate(pages):
        ctx.progress(k / len(pages))
        pg = doc[p]
        try:
            tabs = pg.find_tables(strategy=strategy).tables
        except Exception:
            tabs = []
        rows_per_table = [t.extract() for t in tabs if t.row_count and t.col_count]
        if not rows_per_table:
            lines = [ln for ln in pg.get_text("text").splitlines() if ln.strip()]
            if lines:
                rows_per_table = [[[c for c in re.split(r"\s{2,}|\t", ln.strip())] for ln in lines]]
        for ti, rows in enumerate(rows_per_table):
            n_tables += 1
            if layout == "page":
                ws = wb[f"Page {p + 1}"] if f"Page {p + 1}" in wb.sheetnames else wb.create_sheet(f"Page {p + 1}")
                if ws.max_row > 1:
                    ws.append([])
            elif single is not None:
                ws = single
                ws.append([])
            else:
                ws = wb.create_sheet(f"P{p + 1}-T{ti + 1}"[:31])
            for r_i, row in enumerate(rows):
                ws.append([_cell(c) for c in row])
                if r_i == 0 and layout == "table":
                    for c in ws[ws.max_row]:
                        c.font = Font(bold=True)
    if not wb.sheetnames or n_tables == 0:
        raise ToolError("no_tables", "No tables or text were found in this PDF.")
    for ws in wb.worksheets:
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(60, max(8, width + 2))
    dst = ctx.new_path("converted.xlsx")
    wb.save(str(dst))
    ctx.data["tables"] = n_tables
    ctx.add_output(dst, f"{ctx.stem()}.xlsx")


# ---------------------------------------------------------------------------
@tool("pdf2md", timeout=900, mem_gb=6)
def pdf2md(ctx: Ctx):
    import pymupdf4llm

    src = ctx.paths[0]
    pw = ctx.opts.get("password")
    if pw:
        d = open_pdf(src, pw)
        src = ctx.tmp / "dec.pdf"
        d.save(str(src))
    src = _maybe_ocr(ctx, src)
    doc = open_pdf(src)
    pages = parse_pages(ctx.opt("pages", "all"), len(doc))
    with_images = bool(ctx.opts.get("images"))
    kw = dict(pages=pages)
    img_dir = ctx.tmp / "images"
    if with_images:
        img_dir.mkdir(exist_ok=True)
        kw.update(write_images=True, image_path=str(img_dir), image_format="png", dpi=130)
    try:
        md = pymupdf4llm.to_markdown(doc, **kw)
    except TypeError:
        md = pymupdf4llm.to_markdown(doc, pages=pages)
    md = md if isinstance(md, str) else "\n\n".join(c.get("text", "") for c in md)
    md = re.sub(r"\n{4,}", "\n\n\n", md).strip() + "\n"
    ctx.data["chars"] = len(md)
    if with_images and any(img_dir.iterdir()):
        files = [(Path(ctx.tmp / "doc.md"), f"{ctx.stem()}.md")]
        (ctx.tmp / "doc.md").write_text(md.replace(str(img_dir) + "/", "images/"), encoding="utf-8")
        files += [(f, f"images/{f.name}") for f in sorted(img_dir.iterdir())]
        zip_outputs(ctx, files, f"{ctx.stem()}_markdown.zip")
    else:
        dst = ctx.new_path("converted.md")
        dst.write_text(md, encoding="utf-8")
        ctx.add_output(dst, f"{ctx.stem()}.md")


@tool("pdf2txt", timeout=600)
def pdf2txt(ctx: Ctx):
    doc = pdf_of(ctx, 0)
    pages = parse_pages(ctx.opt("pages", "all"), len(doc))
    text = "\n\n".join(f"--- Page {p + 1} ---\n{doc[p].get_text('text', sort=True).strip()}" for p in pages)
    dst = ctx.new_path("text.txt")
    dst.write_text(text + "\n", encoding="utf-8")
    ctx.add_output(dst, f"{ctx.stem()}.txt")


# ---------------------------------------------------------------------------
def _lo_env(profile: Path) -> dict:
    return {
        "HOME": str(profile.parent),
        "SAL_USE_VCLPLUGIN": "svp",
        "LD_LIBRARY_PATH": str(config.LO_HOME / "program"),
        "TMPDIR": str(profile.parent),
    }


def ensure_lo_profile_template() -> Path | None:
    """Create (once) a locked-down LibreOffice profile used as the per-job seed."""
    tpl = config.DATA / "lo-profile-template"
    if (tpl / "user" / "registrymodifications.xcu").exists():
        return tpl
    soffice = config.LO_HOME / "program" / "soffice"
    if not soffice.exists():
        return None
    tpl.mkdir(parents=True, exist_ok=True)
    work = config.DATA / "lo-warmup"
    work.mkdir(parents=True, exist_ok=True)
    (work / "w.txt").write_text("warmup")
    env = dict(os.environ, **_lo_env(tpl))
    subprocess.run([str(soffice), "--headless", "--norestore", f"-env:UserInstallation=file://{tpl}",
                    "--convert-to", "pdf", "--outdir", str(work), str(work / "w.txt")],
                   env=env, capture_output=True, timeout=180)
    reg = tpl / "user" / "registrymodifications.xcu"
    if not reg.exists():
        return None
    xml = reg.read_text(encoding="utf-8")
    items = (
        '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>\n'
        '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item>\n'
        '<item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item>\n'
        '<item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item>\n')
    reg.write_text(xml.replace("</oor:items>", items + "</oor:items>"), encoding="utf-8")
    shutil.rmtree(work, ignore_errors=True)
    return tpl


def lo_convert(ctx: Ctx, src: Path, outdir: Path, ext="pdf", filt=None) -> Path:
    soffice = config.LO_HOME / "program" / "soffice"
    if not soffice.exists():
        raise ToolError("unavailable", "Office conversion is not available on this server.")
    profile = ctx.tmp / f"lo-profile-{src.stem[:8]}"
    tpl = ensure_lo_profile_template()
    if tpl and not profile.exists():
        shutil.copytree(tpl, profile)
    target = ext if not filt else f"{ext}:{filt}"
    run_cmd([str(soffice), "--headless", "--norestore", "--nolockcheck", "--nodefault", "--nofirststartwizard",
             f"-env:UserInstallation=file://{profile}", "--convert-to", target, "--outdir", str(outdir), str(src)],
            timeout=240, env_extra=_lo_env(profile), cwd=str(ctx.tmp))
    out = outdir / f"{src.stem}.{ext}"
    if not out.exists():
        raise ToolError("convert_failed", "The document could not be converted. It may be corrupt or protected.")
    return out


def _office(ctx: Ctx, kind: str):
    files = []
    for i, p in enumerate(ctx.paths):
        ctx.progress(i / len(ctx.paths), ctx.names[i])
        src = ctx.tmp / f"in{i}{p.suffix.lower()}"
        shutil.copy(p, src)
        if kind == "excel" and ctx.opts.get("fit_width") and src.suffix == ".xlsx":
            _fit_xlsx(src)
        outdir = ctx.tmp / f"o{i}"
        outdir.mkdir()
        res = lo_convert(ctx, src, outdir, "pdf")
        files.append((res, f"{Path(ctx.names[i]).stem}.pdf"))
    if len(files) == 1:
        ctx.add_output(*files[0])
    else:
        zip_outputs(ctx, files, "converted_pdfs.zip")


def _fit_xlsx(path: Path):
    try:
        from openpyxl import load_workbook
        from openpyxl.worksheet.properties import PageSetupProperties

        wb = load_workbook(path)
        for ws in wb.worksheets:
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            ws.page_setup.orientation = "landscape"
        wb.save(path)
    except Exception:
        pass


@tool("word2pdf", accept=("doc", "docx", "odt", "rtf", "txt"), multi=True, max_files=10, timeout=900, mem_gb=None)
def word2pdf(ctx: Ctx):
    _office(ctx, "word")


@tool("ppt2pdf", accept=("ppt", "pptx", "odp"), multi=True, max_files=10, timeout=900, mem_gb=None)
def ppt2pdf(ctx: Ctx):
    _office(ctx, "ppt")


@tool("excel2pdf", accept=("xls", "xlsx", "ods", "csv"), multi=True, max_files=10, timeout=900, mem_gb=None)
def excel2pdf(ctx: Ctx):
    _office(ctx, "excel")


# ---------------------------------------------------------------------------
def _chromium_exe() -> str | None:
    base = Path(os.environ.get("PLAYWRIGHT_BROWSERS_PATH") or Path.home() / ".cache" / "ms-playwright")
    cands = sorted(base.glob("chromium_headless_shell-*/chrome-*/headless_shell"), reverse=True)
    cands += sorted(base.glob("chromium-*/chrome-linux*/chrome"), reverse=True)
    return str(cands[0]) if cands else None


@tool("html2pdf", accept=("html", "htm"), min_files=0, timeout=240, mem_gb=None)
def html2pdf(ctx: Ctx):
    from playwright.sync_api import sync_playwright

    url = str(ctx.opt("url", "")).strip()
    html = ctx.opt("html", "")
    if not url and not html and not ctx.paths:
        raise ToolError("nothing", "Enter a web address, paste HTML or upload an .html file.")
    if url:
        url = netguard.check_url(url)
    exe = _chromium_exe()
    if not exe:
        raise ToolError("unavailable", "The browser engine is not available on this server.")
    size = ctx.opt("page_size", "A4")
    landscape = ctx.opt("orientation", "portrait") == "landscape"
    margin = f"{float(ctx.opt('margin', 10))}mm"
    width = max(320, min(2400, int(ctx.opt("viewport", 1280))))
    js_on = bool(ctx.opts.get("javascript", bool(url)))
    single = bool(ctx.opts.get("single_page"))
    media = "screen" if ctx.opt("media", "print") == "screen" else "print"
    with sync_playwright() as pw:
        browser = pw.chromium.launch(
            executable_path=exe, headless=True,
            args=["--no-sandbox", "--disable-gpu", "--disable-dev-shm-usage", "--disable-extensions",
                  "--mute-audio", "--no-first-run", "--disable-background-networking"],
            env={"PATH": os.environ.get("PATH", ""), "HOME": str(ctx.tmp), "LANG": "C.UTF-8"})
        try:
            context = browser.new_context(viewport={"width": width, "height": 900}, java_script_enabled=js_on,
                                          accept_downloads=False, service_workers="block")

            def guard(route):
                if netguard.request_allowed(route.request.url):
                    route.continue_()
                else:
                    route.abort("blockedbyclient")

            context.route("**/*", guard)
            context.add_init_script("window.WebSocket = function(){ throw new Error('blocked'); };")  # no WS to internal hosts
            page = context.new_page()
            page.set_default_timeout(30000)
            try:
                if url:
                    page.goto(url, wait_until="load", timeout=40000)
                    try:
                        page.wait_for_load_state("networkidle", timeout=8000)
                    except Exception:
                        pass
                elif html:
                    page.set_content(str(html), wait_until="load")
                else:
                    page.set_content(ctx.paths[0].read_text("utf-8", "replace"), wait_until="load")
            except Exception as e:
                raise ToolError("load_failed", "The page could not be loaded: " + str(e).split("\n")[0][:160])
            if js_on:
                try:  # trigger lazy-loaded images
                    page.evaluate("""async()=>{for(let y=0;y<document.body.scrollHeight&&y<30000;y+=700){window.scrollTo(0,y);await new Promise(r=>setTimeout(r,60));}window.scrollTo(0,0);}""")
                except Exception:
                    pass
            wait = max(0.0, min(10.0, float(ctx.opt("wait", 0))))
            if wait:
                page.wait_for_timeout(int(wait * 1000))
            page.emulate_media(media=media)
            kw = dict(print_background=bool(ctx.opts.get("background", True)), landscape=landscape,
                      margin=dict(top=margin, bottom=margin, left=margin, right=margin))
            if single:
                h = min(14000, int(page.evaluate("Math.max(document.documentElement.scrollHeight, document.body.scrollHeight)")) + 4)
                kw = dict(print_background=kw["print_background"], width=f"{width}px", height=f"{h}px",
                          margin=dict(top="0", bottom="0", left="0", right="0"), page_ranges="1")
            else:
                kw["format"] = size if size in ("A3", "A4", "A5", "Letter", "Legal", "Tabloid") else "A4"
            dst = ctx.new_path("page.pdf")
            page.pdf(path=str(dst), **kw)
        finally:
            browser.close()
    name = "page"
    if url:
        from urllib.parse import urlsplit
        name = safe_stem(urlsplit(url).hostname or "page")
    elif ctx.paths:
        name = ctx.stem()
    ctx.add_output(dst, f"{name}.pdf")
